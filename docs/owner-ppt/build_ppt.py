#!/usr/bin/env python3
"""Owner onboarding PPT with Web + iOS + Android screenshots."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
WEB = SHOTS / "web"
IOS = SHOTS / "mobile"  # iPhone captures
ANDROID = SHOTS / "android"
LEGACY = SHOTS  # fallback if platform folders missing
OUT = ROOT / "BuildFlow_Owner_Construction_ERP_Guide.pptx"
ICON = Path("/home/prasanna/work/BuildFlow/docs/owner-ppt/assets/buildflow-logo.png")
if not ICON.exists():
    ICON = Path("/home/prasanna/work/BuildFlow/apps/mobile/assets/icon.png")


PRIMARY = RGBColor(0x1E, 0x3A, 0x5F)
ACCENT = RGBColor(0xF5, 0x9E, 0x0B)
TEXT = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SURFACE = RGBColor(0xF8, 0xFA, 0xFC)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, *, size=18, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_text(shape, text, *, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)


def add_para(tf, text, *, size=16, bold=False, color=TEXT, space_before=6):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def paint_bg(slide, color=WHITE):
    fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    fill.line.fill.background()
    fill.fill.solid()
    fill.fill.fore_color.rgb = color
    spTree = slide.shapes._spTree
    sp = fill._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def accent_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.1))
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT


def footer(slide, page_no: int, total: int):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(7.12), Inches(10), Inches(0.28))
    add_text(box, "BuildFlow Construction ERP  ·  Android · iOS · Web  ·  Owner Quick-Start", size=10, color=MUTED)
    num = slide.shapes.add_textbox(Inches(11.5), Inches(7.12), Inches(1.4), Inches(0.28))
    add_text(num, f"{page_no} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def title_block(slide, title: str, subtitle: str | None = None):
    t = slide.shapes.add_textbox(Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.45))
    add_text(t, title, size=26, bold=True, color=PRIMARY)
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.5), Inches(0.72), Inches(12.3), Inches(0.35))
        add_text(s, subtitle, size=13, color=MUTED)


def resolve_shot(*candidates: Path) -> Path | None:
    for p in candidates:
        if p.exists() and p.stat().st_size > 20000:
            return p
    return None


def add_pic(slide, path: Path | None, left, top, width, height=None):
    if path is None:
        missing = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height or Inches(3.2)
        )
        missing.fill.solid()
        missing.fill.fore_color.rgb = SURFACE
        missing.line.color.rgb = BORDER
        add_text(missing, "Screenshot pending", size=11, color=MUTED, align=PP_ALIGN.CENTER)
        return
    if height:
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    else:
        slide.shapes.add_picture(str(path), left, top, width=width)


def platform_badge(slide, left, top, label: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(1.35), Inches(0.28))
    badge.fill.solid()
    badge.fill.fore_color.rgb = PRIMARY
    badge.line.fill.background()
    add_text(badge, label, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)


def bullet_card(slide, left, top, width, height, title, bullets):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = BORDER
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    set_run(run, size=14, bold=True, color=PRIMARY)
    for b in bullets:
        add_para(tf, f"•  {b}", size=12, color=TEXT, space_before=7)


def tri_platform_slide(
    prs,
    *,
    title: str,
    subtitle: str,
    shot_name: str,
    legacy_names: list[str] | None = None,
    tips: list[str],
):
    """Web (wide) + phone (iOS) + phone (Android)."""
    s = blank_slide(prs)
    paint_bg(s)
    accent_bar(s)
    title_block(s, title, subtitle)

    legacy = legacy_names or []
    web = resolve_shot(WEB / f"{shot_name}.png", *[LEGACY / n for n in legacy])
    ios = resolve_shot(IOS / f"{shot_name}.png", IOS / f"ios-{shot_name}.png")
    android = resolve_shot(ANDROID / f"{shot_name}.png")

    # Web large left
    platform_badge(s, Inches(0.5), Inches(1.15), "WEB")
    add_pic(s, web, Inches(0.5), Inches(1.5), Inches(7.4), Inches(4.55))

    # Phones stacked/side on right
    platform_badge(s, Inches(8.15), Inches(1.15), "iOS")
    add_pic(s, ios, Inches(8.15), Inches(1.5), Inches(2.15), Inches(4.55))

    platform_badge(s, Inches(10.55), Inches(1.15), "ANDROID")
    add_pic(s, android, Inches(10.55), Inches(1.5), Inches(2.25), Inches(4.55))

    # Tip strip
    tip = s.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(12.3), Inches(0.75))
    tf = tip.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Same account · same data  ·  " + "  ·  ".join(tips)
    set_run(run, size=12, color=MUTED)
    return s


def web_focus_slide(prs, title, subtitle, shot_name, legacy, tips):
    s = blank_slide(prs)
    paint_bg(s)
    accent_bar(s)
    title_block(s, title, subtitle)
    web = resolve_shot(WEB / f"{shot_name}.png", *[LEGACY / n for n in legacy])
    ios = resolve_shot(IOS / f"{shot_name}.png")
    add_pic(s, web, Inches(0.4), Inches(1.25), Inches(8.6), Inches(5.35))
    platform_badge(s, Inches(0.4), Inches(1.0), "WEB")
    if ios:
        platform_badge(s, Inches(9.2), Inches(1.0), "PHONE")
        add_pic(s, ios, Inches(9.2), Inches(1.35), Inches(2.35), Inches(5.1))
    bullet_card(s, Inches(11.7) if ios else Inches(9.2), Inches(1.35), Inches(1.4) if ios else Inches(3.7), Inches(5.1) if not ios else Inches(0.01), "", [])
    if not ios:
        bullet_card(s, Inches(9.2), Inches(1.35), Inches(3.7), Inches(5.15), "Owner tip", tips)
    else:
        # tips under phones area as narrow card won't fit — use bottom strip
        tip = s.shapes.add_textbox(Inches(9.2), Inches(6.55), Inches(3.7), Inches(0.4))
        add_text(tip, " · ".join(tips[:2]), size=10, color=MUTED)
    return s


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 Title
    s = blank_slide(prs)
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(5.5), SLIDE_H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PRIMARY
    panel.line.fill.background()
    if ICON.exists():
        s.shapes.add_picture(str(ICON), Inches(0.75), Inches(1.35), width=Inches(0.9))
    brand = s.shapes.add_textbox(Inches(0.75), Inches(2.45), Inches(4.3), Inches(0.7))
    add_text(brand, "BuildFlow", size=40, bold=True, color=WHITE)
    sub = s.shapes.add_textbox(Inches(0.75), Inches(3.15), Inches(4.3), Inches(0.4))
    add_text(sub, "Construction ERP", size=20, color=ACCENT)
    plat = s.shapes.add_textbox(Inches(0.75), Inches(3.75), Inches(4.3), Inches(0.4))
    add_text(plat, "Android  ·  iOS  ·  Web", size=16, bold=True, color=WHITE)
    tag = s.shapes.add_textbox(Inches(0.75), Inches(4.4), Inches(4.3), Inches(1.4))
    tf = tag.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Owner Quick-Start Guide"
    set_run(run, size=22, bold=True, color=WHITE)
    add_para(tf, "One company account. Use it on phone at site and on desktop in the office.", size=13, color=RGBColor(0xCB, 0xD5, 0xE1), space_before=10)

    right = s.shapes.add_textbox(Inches(6.1), Inches(1.7), Inches(6.5), Inches(4.2))
    tf = right.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "Ready for owners who run sites & boards"
    set_run(run, size=24, bold=True, color=PRIMARY)
    for line in [
        "Approve estimates & variations from anywhere",
        "See budgets and receivables on Owner Home",
        "Invite PM, site, store, QC, accounts in minutes",
        "Same live data on Android, iOS, and Web",
    ]:
        add_para(tf, f"✓  {line}", size=15, color=TEXT, space_before=12)
    foot = s.shapes.add_textbox(Inches(6.1), Inches(6.5), Inches(6.5), Inches(0.4))
    add_text(foot, "Screenshots: live Reddy Constructions demo  ·  BuildFlow v2.0", size=11, color=MUTED)

    # 2 Platforms overview
    s = blank_slide(prs)
    paint_bg(s)
    accent_bar(s)
    title_block(s, "One product — three surfaces", "Owners and teams use the same Construction ERP on phone and desktop")
    cards = [
        ("Android", "Phone & tablet at site, store, weighbridge"),
        ("iOS", "iPhone / iPad for PM & Owner on the move"),
        ("Web", "Desktop for estimates, accounting, approvals"),
    ]
    for i, (title, desc) in enumerate(cards):
        left = Inches(0.55 + i * 4.2)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.6), Inches(3.95), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE
        card.line.color.rgb = BORDER
        t = s.shapes.add_textbox(left + Inches(0.25), Inches(1.9), Inches(3.4), Inches(0.5))
        add_text(t, title, size=22, bold=True, color=PRIMARY)
        d = s.shapes.add_textbox(left + Inches(0.25), Inches(2.55), Inches(3.4), Inches(0.9))
        tf = d.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = desc
        set_run(run, size=14, color=TEXT)
    note = s.shapes.add_textbox(Inches(0.55), Inches(4.2), Inches(12.2), Inches(2.2))
    tf = note.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "How to open it"
    set_run(run, size=16, bold=True, color=PRIMARY)
    for line in [
        "Web: open your BuildFlow company URL in Chrome / Edge / Safari",
        "Android: install the BuildFlow app (or open the same URL in Chrome)",
        "iOS: install the BuildFlow app (or open the same URL in Safari / Add to Home Screen)",
        "Login is always email/mobile + OTP — no passwords to share on WhatsApp",
    ]:
        add_para(tf, f"•  {line}", size=14, color=TEXT, space_before=10)

    # 3 Agenda
    s = blank_slide(prs)
    paint_bg(s)
    accent_bar(s)
    title_block(s, "What you will learn", "Practical path so your company starts using BuildFlow this week")
    items = [
        ("01", "Sign in (OTP)"),
        ("02", "Owner Home"),
        ("03", "Projects"),
        ("04", "Estimates → BOQ"),
        ("05", "Procurement"),
        ("06", "Accounting"),
        ("07", "Invite team"),
        ("08", "Day-1 checklist"),
    ]
    for i, (num, label) in enumerate(items):
        col, row = i % 4, i // 4
        left = Inches(0.55 + col * 3.15)
        top = Inches(1.6 + row * 2.3)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.95), Inches(1.9))
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE
        card.line.color.rgb = BORDER
        nbox = s.shapes.add_textbox(left + Inches(0.2), top + Inches(0.4), Inches(2.5), Inches(0.45))
        add_text(nbox, num, size=26, bold=True, color=ACCENT)
        lbox = s.shapes.add_textbox(left + Inches(0.2), top + Inches(1.0), Inches(2.5), Inches(0.55))
        add_text(lbox, label, size=15, bold=True, color=PRIMARY)

    # Content slides — tri platform
    tri_platform_slide(
        prs,
        title="1. Sign in — email or mobile + OTP",
        subtitle="Identical login on Android, iOS, and Web. Owners and field staff use the same secure flow.",
        shot_name="01-login",
        legacy_names=["01-login.png", "01b-login-filled.png"],
        tips=["Send OTP", "Enter 6-digit code", "Land on Owner Home"],
    )
    tri_platform_slide(
        prs,
        title="2. Owner Home — company pulse",
        subtitle="Active projects, receivables, progress, budget burn, team activity — on phone or desktop.",
        shot_name="02-home",
        legacy_names=["03-dashboard.png", "02-after-login.png"],
        tips=["KPIs at a glance", "Quick actions", "Approve from anywhere"],
    )
    tri_platform_slide(
        prs,
        title="3. Projects — your portfolio",
        subtitle="Create jobs, set budgets, open any site. Desktop for planning; phone for quick checks.",
        shot_name="03-projects",
        legacy_names=["04-projects.png"],
        tips=["+ New Project", "Filter Active / Planning", "Open NH-45 sample"],
    )
    tri_platform_slide(
        prs,
        title="4. Inside a project — Overview",
        subtitle="Budget, committed spend, schedule health. Money stays Owner/PM visible.",
        shot_name="10-project",
        legacy_names=["13-project-overview.png"],
        tips=["Overview KPIs", "Tabs for BOQ & more", "Setup checklist"],
    )
    tri_platform_slide(
        prs,
        title="5. Estimates → BOQ",
        subtitle="Approve commercially on web; track quantities on site with the phone.",
        shot_name="11-boq",
        legacy_names=["14-tab-boq.png"],
        tips=["Owner approves estimate", "Convert to BOQ", "Site measures qty"],
    )
    tri_platform_slide(
        prs,
        title="6. Procurement — indents, PO, GRN",
        subtitle="PM/Store raise indents; weighbridge/store receive on mobile at the gate.",
        shot_name="12-procurement",
        legacy_names=["16-tab-procurement.png"],
        tips=["Indent from BOQ", "Approve PO", "GRN on phone"],
    )
    tri_platform_slide(
        prs,
        title="7. Variations — Owner approval",
        subtitle="PM/DPM draft change orders; only Owner signs off cost & schedule impact.",
        shot_name="13-variations",
        legacy_names=["17-tab-variations.png"],
        tips=["Review impact", "Approve / reject", "Convert VO → BOQ"],
    )
    tri_platform_slide(
        prs,
        title="8. Accounting — invoices & bills",
        subtitle="Desktop for finance ops; phone for quick status. Accountant can approve bills.",
        shot_name="04-accounting",
        legacy_names=["05-accounting.png"],
        tips=["Invoices & bills", "Record payments", "Tally / GST export"],
    )
    tri_platform_slide(
        prs,
        title="9. Planning & delivery",
        subtitle="CPM-style planning on web; progress checks on mobile.",
        shot_name="07-planning",
        legacy_names=["09-planning.png"],
        tips=["Tasks & schedule", "Progress %", "Overdue visibility"],
    )
    tri_platform_slide(
        prs,
        title="10. Reports Hub",
        subtitle="Board packs and PDF downloads — best on web; share links from phone.",
        shot_name="05-reports-hub",
        legacy_names=["07-reports-hub.png"],
        tips=["Progress PDFs", "P&L / EVA", "GST & TDS"],
    )
    tri_platform_slide(
        prs,
        title="11. Invite your team",
        subtitle="Settings → Users on web (recommended). Roles: PM, Site, Store, QC, Accountant…",
        shot_name="09-settings-users",
        legacy_names=["11-settings-users.png"],
        tips=["Invite by email/phone", "Pick role", "OTP onboarding"],
    )

    # Roles
    s = blank_slide(prs)
    paint_bg(s)
    accent_bar(s)
    title_block(s, "Who does what", "Share with leadership — works the same on every device")
    roles = [
        ("Owner", "Approve estimates & VOs · finance · invite users · settings"),
        ("PM / DPM", "Run projects · estimates · procurement · draft variations"),
        ("Site / QC", "Daily reports · measure · drawings · snags (no money)"),
        ("Store / Weighbridge", "Indents · GRN · stock on phone at gate"),
        ("Accountant", "Invoices · bills · payments · Tally / GST"),
    ]
    for i, (role, desc) in enumerate(roles):
        top = Inches(1.4 + i * 1.0)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), top, Inches(12.3), Inches(0.88))
        card.fill.solid()
        card.fill.fore_color.rgb = SURFACE if i % 2 == 0 else WHITE
        card.line.color.rgb = BORDER
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), top, Inches(0.12), Inches(0.88))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT if i == 0 else PRIMARY
        bar.line.fill.background()
        r = s.shapes.add_textbox(Inches(0.85), top + Inches(0.2), Inches(2.8), Inches(0.5))
        add_text(r, role, size=17, bold=True, color=PRIMARY)
        d = s.shapes.add_textbox(Inches(3.7), top + Inches(0.25), Inches(8.8), Inches(0.5))
        add_text(d, desc, size=14, color=TEXT)

    # Day-1
    s = blank_slide(prs)
    paint_bg(s)
    accent_bar(s)
    title_block(s, "Day-1 checklist", "Use phone + web together — start today")
    bullet_card(
        s,
        Inches(0.5),
        Inches(1.4),
        Inches(6.05),
        Inches(5.2),
        "Morning (Web or iPad)",
        [
            "Sign in as Owner (OTP)",
            "Review Owner Home KPIs",
            "Create / open your real project",
            "Set budget, client, dates",
            "Invite PM + Site + Accountant",
            "Optional: walk NH-45 sample",
        ],
    )
    bullet_card(
        s,
        Inches(6.8),
        Inches(1.4),
        Inches(6.05),
        Inches(5.2),
        "Afternoon (Phone + Web)",
        [
            "PM drafts first estimate (Web)",
            "You approve on phone or desktop",
            "Convert to BOQ",
            "First indent → GRN on Android/iOS",
            "First vendor bill → Accountant",
            "Download one PDF from Reports Hub",
        ],
    )

    # Closing
    s = blank_slide(prs)
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PRIMARY
    panel.line.fill.background()
    if ICON.exists():
        s.shapes.add_picture(str(ICON), Inches(6.15), Inches(1.2), width=Inches(1.0))
    t = s.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(10.9), Inches(1))
    add_text(t, "Android · iOS · Web — you’re ready", size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    sub = s.shapes.add_textbox(Inches(2.0), Inches(3.6), Inches(9.3), Inches(1.3))
    tf = sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Owner approves. PM delivers. Site executes. Store receives. Accounts closes the books — one BuildFlow company everywhere."
    set_run(run, size=15, color=RGBColor(0xCB, 0xD5, 0xE1))
    tip = s.shapes.add_textbox(Inches(2.2), Inches(5.3), Inches(8.9), Inches(0.9))
    tf = tip.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Install on phones this week. Run the first live estimate → BOQ → indent cycle on one project."
    set_run(run, size=14, color=ACCENT)

    total = len(prs.slides)
    for idx, slide in enumerate(prs.slides):
        if idx in (0, total - 1):
            continue
        footer(slide, idx + 1, total)

    prs.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size // 1024} KB), {total} slides")


if __name__ == "__main__":
    build()
