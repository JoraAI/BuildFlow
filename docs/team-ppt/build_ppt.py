#!/usr/bin/env python3
"""Build BuildFlow internal team walkthrough PPTX from captured screenshots."""
from __future__ import annotations

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
SHOT = os.path.join(OUT_DIR, "screenshots")
OUT_PPTX = os.path.join(OUT_DIR, "BuildFlow_Team_Walkthrough.pptx")

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
slides_meta: list = []


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def set_run(p, text, size=18, bold=False, color=SLATE, font="Calibri"):
    p.clear()
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return r


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=SLATE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p, text, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, l, t, w, h, items, size=15, color=SLATE):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        set_run(p, "•  " + item, size=size, color=color)
    return box


def footer(slide, page, total):
    add_text(
        slide,
        Inches(0.4),
        Inches(7.15),
        Inches(8),
        Inches(0.3),
        "BuildFlow v2.0  ·  Internal team walkthrough  ·  Confidential",
        size=10,
        color=MUTED,
    )
    add_text(
        slide,
        Inches(11.5),
        Inches(7.15),
        Inches(1.5),
        Inches(0.3),
        f"{page} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def section_bar(slide, title, subtitle=None):
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.95), NAVY)
    add_text(slide, Inches(0.45), Inches(0.18), Inches(10), Inches(0.4), title, size=26, bold=True, color=WHITE)
    if subtitle:
        add_text(
            slide,
            Inches(0.45),
            Inches(0.55),
            Inches(11),
            Inches(0.3),
            subtitle,
            size=12,
            color=RGBColor(0xCB, 0xD5, 0xE1),
        )
    add_rect(slide, Inches(0), Inches(0.95), Inches(13.333), Inches(0.08), AMBER)


def img(slide, name, l, t, w, h=None):
    path = os.path.join(SHOT, name)
    if not os.path.exists(path):
        add_text(slide, l, t, w, Inches(0.4), f"[missing {name}]", size=12, color=MUTED)
        return None
    if h:
        return slide.shapes.add_picture(path, l, t, width=w, height=h)
    return slide.shapes.add_picture(path, l, t, width=w)


def dual_viewport(slide, web_name, mobile_name, caption_web, caption_mobile):
    img(slide, web_name, Inches(0.4), Inches(1.25), Inches(8.4))
    add_text(slide, Inches(0.4), Inches(6.55), Inches(8.4), Inches(0.3), caption_web, size=11, color=MUTED)
    img(slide, mobile_name, Inches(9.1), Inches(1.25), Inches(3.7))
    add_text(
        slide,
        Inches(9.1),
        Inches(6.55),
        Inches(3.7),
        Inches(0.3),
        caption_mobile,
        size=11,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def new_slide():
    s = prs.slides.add_slide(BLANK)
    slides_meta.append(s)
    return s


# ── 1 Title ──────────────────────────────────────────────────────────
s = new_slide()
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
add_rect(s, Inches(0), Inches(6.6), Inches(13.333), Inches(0.12), AMBER)
add_text(s, Inches(0.8), Inches(2.0), Inches(11.5), Inches(0.8), "BuildFlow", size=54, bold=True, color=WHITE)
add_text(
    s,
    Inches(0.8),
    Inches(2.85),
    Inches(11.5),
    Inches(0.5),
    "Product walkthrough for the team",
    size=28,
    color=AMBER,
)
add_text(
    s,
    Inches(0.8),
    Inches(3.6),
    Inches(11),
    Inches(0.8),
    "Construction ERP  ·  Inventory / Stock  ·  Platform Admin\n"
    "All major functions, roles, and why it matters — with live web & mobile screens.",
    size=16,
    color=RGBColor(0xE2, 0xE8, 0xF0),
)
add_text(
    s,
    Inches(0.8),
    Inches(5.5),
    Inches(11),
    Inches(0.5),
    "Demo logins used: Platform Admin + Company Owner (full access)  ·  OTP seed: 111111",
    size=13,
    color=RGBColor(0x94, 0xA3, 0xB8),
)

# ── 2 Agenda ─────────────────────────────────────────────────────────
s = new_slide()
section_bar(s, "Agenda", "What this deck covers")
add_bullets(
    s,
    Inches(0.6),
    Inches(1.4),
    Inches(6),
    Inches(5),
    [
        "What BuildFlow is (two products, one platform)",
        "Why teams adopt it — product strengths",
        "Roles across Construction, Inventory & Platform",
        "How to sign in (admin / owner demos)",
        "Construction ERP — every major module",
        "Inventory / Stock — every major module",
        "Platform Admin console",
        "Suggested team walkthrough paths",
    ],
    size=18,
)
add_rect(s, Inches(7.4), Inches(1.5), Inches(5.3), Inches(4.8), LIGHT)
add_text(s, Inches(7.7), Inches(1.75), Inches(4.8), Inches(0.4), "Screenshot convention", size=16, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(7.7),
    Inches(2.3),
    Inches(4.8),
    Inches(3.5),
    [
        "Left / wide = Desktop web (1440×900)",
        "Right / tall = Mobile viewport (390×844)",
        "Captured from running local app",
        "Owner role for full module access",
        "Platform Admin for tenant ops",
    ],
    size=14,
)

# ── 3 What is BuildFlow ──────────────────────────────────────────────
s = new_slide()
section_bar(s, "What is BuildFlow?", "One codebase · two product modes · role-aware UX")
add_rect(s, Inches(0.5), Inches(1.4), Inches(5.9), Inches(5.2), LIGHT)
add_text(s, Inches(0.8), Inches(1.65), Inches(5.3), Inches(0.4), "Construction ERP", size=22, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(0.8),
    Inches(2.3),
    Inches(5.3),
    Inches(4),
    [
        "Projects, BOQ & estimation",
        "Proposals / tendering",
        "CPM planning & schedules",
        "Daily site reports & attendance",
        "Procurement, stock, subcontractors",
        "Accounting (invoices, bills, Tally)",
        "AI assistant + notifications",
        "Company settings, users, permissions",
    ],
    size=15,
)
add_rect(s, Inches(6.9), Inches(1.4), Inches(5.9), Inches(5.2), LIGHT)
add_text(s, Inches(7.2), Inches(1.65), Inches(5.3), Inches(0.4), "Inventory / Stock", size=22, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(7.2),
    Inches(2.3),
    Inches(5.3),
    Inches(4),
    [
        "Stock home with scan & checkout",
        "Materials / SKU catalog",
        "Parties (customers & vendors)",
        "Quotes → Sales → Invoices",
        "Warehouse transfers & GRN",
        "Procurement / purchase bills",
        "Reports + GST-oriented flows",
        "Profiles: retail, wholesale, materials, events…",
    ],
    size=15,
)

# ── 4 Pros ───────────────────────────────────────────────────────────
s = new_slide()
section_bar(s, "Why BuildFlow — product strengths", "Talking points for the team")
pros = [
    ("Field + boardroom", "Same app on phone and desktop — supervisors on site, owners in the office."),
    ("Role-based access", "Permissions per role; Owners can customize. Site roles hide amounts when needed."),
    ("Two products, shared platform", "Construction ERP and Inventory tenants on one stack; Platform Admin oversees all."),
    ("India-ready ops", "₹ currency, GST-oriented invoices/bills, Tally export, OTP login (email/SMS)."),
    ("End-to-end construction", "Estimate → BOQ → procure → measure → invoice — not disconnected tools."),
    ("Inventory verticals", "Material supplier, retail/kirana FEFO, events/lighting wording, multi-warehouse."),
    ("AI assistant", "In-app assistant FAB with permission-aware tools for faster ops."),
    ("Audit & support", "Audit log, tickets, escalations to Platform — built for multi-tenant SaaS."),
]
for i, (title, body) in enumerate(pros):
    col = i % 2
    row = i // 2
    x = Inches(0.4 + col * 6.45)
    y = Inches(1.25 + row * 1.4)
    add_rect(s, x, y, Inches(6.2), Inches(1.25), LIGHT)
    add_rect(s, x, y, Inches(0.12), Inches(1.25), AMBER)
    add_text(s, x + Inches(0.3), y + Inches(0.15), Inches(5.7), Inches(0.35), title, size=15, bold=True, color=NAVY)
    add_text(s, x + Inches(0.3), y + Inches(0.5), Inches(5.7), Inches(0.6), body, size=12, color=SLATE)

# ── 5 Roles overview ─────────────────────────────────────────────────
s = new_slide()
section_bar(s, "Roles at a glance", "Construction · Inventory · Platform")
add_text(s, Inches(0.5), Inches(1.25), Inches(4), Inches(0.35), "Construction ERP", size=16, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(0.5),
    Inches(1.7),
    Inches(4),
    Inches(5),
    [
        "Owner / MD — full access",
        "Project Manager (PM)",
        "Deputy PM (DPM)",
        "Senior QC Engineer",
        "Mechanical Manager",
        "Store Incharge",
        "WeighBridge Incharge",
        "Site Supervisor",
        "Accountant",
    ],
    size=14,
)
add_text(s, Inches(4.8), Inches(1.25), Inches(4), Inches(0.35), "Inventory product", size=16, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(4.8),
    Inches(1.7),
    Inches(4),
    Inches(2.5),
    [
        "Owner — full inventory + settings",
        "Inventory Manager — day-to-day stock, sales, procurement",
    ],
    size=14,
)
add_text(s, Inches(4.8), Inches(3.6), Inches(4), Inches(0.35), "Platform", size=16, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(4.8),
    Inches(4.05),
    Inches(4),
    Inches(2),
    [
        "Platform Admin — tenants, subscriptions, escalated tickets",
        "Separate login: /platform/login",
    ],
    size=14,
)
add_rect(s, Inches(9.0), Inches(1.25), Inches(3.9), Inches(5.4), LIGHT)
add_text(s, Inches(9.25), Inches(1.5), Inches(3.4), Inches(0.4), "Design principle", size=14, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(9.25),
    Inches(2.1),
    Inches(3.4),
    Inches(4.2),
    [
        "OWNER gets every permission",
        "Field roles focus on reports / stock / QC",
        "Accountant owns money flows + Tally",
        "Site Supervisor can hide financials",
        "Company can customize role permissions",
    ],
    size=13,
)

# ── 6 Demo credentials ───────────────────────────────────────────────
s = new_slide()
section_bar(s, "Demo credentials (seed)", "Use these for local walkthroughs — password/OTP as noted")
add_rect(s, Inches(0.4), Inches(1.25), Inches(12.5), Inches(1.8), LIGHT)
add_text(
    s,
    Inches(0.7),
    Inches(1.4),
    Inches(12),
    Inches(0.35),
    "Platform Admin (used for platform screens in this deck)",
    size=15,
    bold=True,
    color=NAVY,
)
add_text(
    s,
    Inches(0.7),
    Inches(1.9),
    Inches(12),
    Inches(0.9),
    "URL:  /platform/login\nEmail:  admin@buildflow.com\nPassword:  Admin@1234",
    size=14,
    color=SLATE,
)
add_rect(s, Inches(0.4), Inches(3.3), Inches(6.1), Inches(3.3), LIGHT)
add_text(s, Inches(0.7), Inches(3.5), Inches(5.6), Inches(0.35), "Construction (Owner)", size=15, bold=True, color=NAVY)
add_text(
    s,
    Inches(0.7),
    Inches(4.0),
    Inches(5.6),
    Inches(2.2),
    "Company: Reddy Constructions Pvt Ltd\n"
    "Email: owner@reddyconst.com\n"
    "OTP: 111111 (seed master code)\n\n"
    "Also: pm@ · dpm@ · qc@ · store@ ·\n"
    "site@ · accounts@ · weighbridge@ ·\n"
    "mechanical@reddyconst.com  (same OTP)",
    size=13,
    color=SLATE,
)
add_rect(s, Inches(6.8), Inches(3.3), Inches(6.1), Inches(3.3), LIGHT)
add_text(s, Inches(7.1), Inches(3.5), Inches(5.6), Inches(0.35), "Inventory (Owner)", size=15, bold=True, color=NAVY)
add_text(
    s,
    Inches(7.1),
    Inches(4.0),
    Inches(5.6),
    Inches(2.2),
    "Materials demo (richest):\n"
    "owner@hydmaterials.com  ·  OTP 111111\n\n"
    "Also: kirana-demo · luminalighting ·\n"
    "cityhardware · deccanwholesale ·\n"
    "southdistro · apextrading · forgeequip ·\n"
    "generalstore  (see DEMO_CREDENTIALS.txt)",
    size=13,
    color=SLATE,
)

# ── 7 Login screens ──────────────────────────────────────────────────
s = new_slide()
section_bar(s, "Sign-in experience", "Company OTP login + Platform password login")
dual_viewport(
    s,
    "web_login.png",
    "mobile_login.png",
    "Web — company login (email/mobile + OTP)",
    "Mobile — same flow",
)

# ── 8 Platform ───────────────────────────────────────────────────────
s = new_slide()
section_bar(s, "Platform Admin", "Tenant oversight — admin@buildflow.com")
dual_viewport(
    s,
    "web_platform_home.png",
    "mobile_platform_home.png",
    "Web — companies count, tickets, navigation",
    "Mobile viewport",
)

s = new_slide()
section_bar(s, "Platform — Companies & tickets", "Search tenants, subscriptions, escalations")
img(s, "web_platform_companies.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_platform_tickets.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Companies: manage subscriptions & legal fields   ·   Tickets: review owner escalations",
    size=12,
    color=MUTED,
)

# ── Construction modules ─────────────────────────────────────────────
s = new_slide()
section_bar(s, "Construction — Owner Dashboard", "Projects health, cash, budget burn, quick actions")
dual_viewport(
    s,
    "web_construction_dashboard.png",
    "mobile_construction_dashboard.png",
    "Web — Owner Dashboard (Reddy Constructions)",
    "Mobile — condensed home",
)

s = new_slide()
section_bar(s, "Construction — Projects", "Project list, create, detail (BOQ, procurement, COs…)")
dual_viewport(
    s,
    "web_construction_projects.png",
    "mobile_construction_projects.png",
    "Web — Projects workspace",
    "Mobile — project list",
)

s = new_slide()
section_bar(s, "Construction — Proposals & Estimation", "Tenders, estimates, rate analysis, BOQ conversion")
img(s, "web_construction_proposals.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_construction_estimation.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Proposals  ·  Right: Estimation workspace",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Construction — Planning & Site Reports", "Schedules + daily operational reporting")
img(s, "web_construction_planning.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_construction_reports.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Planning  ·  Right: Reports (site ops)",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Construction — Accounting & Reports Hub", "Invoices, bills, payments · downloadable finance reports")
img(s, "web_construction_accounting.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_construction_reports_hub.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Accounting  ·  Right: Reports Hub",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Construction — Assistant & Alerts", "Permission-aware AI assistant + in-app notifications")
img(s, "web_construction_assistant.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_construction_notifications.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: BuildFlow Assistant (overlay on dashboard)  ·  Right: Alerts / notifications",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Construction — Admin settings", "Company, users/invites, role permissions")
img(s, "web_construction_settings.png", Inches(0.3), Inches(1.2), Inches(4.2))
img(s, "web_construction_users.png", Inches(4.55), Inches(1.2), Inches(4.2))
img(s, "web_construction_permissions.png", Inches(8.8), Inches(1.2), Inches(4.2))
add_text(
    s,
    Inches(0.3),
    Inches(6.55),
    Inches(12.7),
    Inches(0.35),
    "Settings home  ·  Users & invites  ·  Role permissions",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Construction — function map", "Everything under the ERP product mode")
cols = [
    ("Workspace", ["Home / Owner Dashboard", "Projects (create & detail)", "Planning / CPM"]),
    (
        "Operations",
        [
            "Proposals & tenders",
            "Estimation & rate analysis",
            "Site reports & attendance",
            "Drawings / snags / RFIs (in project)",
        ],
    ),
    (
        "Finance",
        [
            "Accounting (AR/AP)",
            "Reports Hub downloads",
            "Petty cash / change orders (project)",
            "Tally export",
        ],
    ),
    (
        "Admin & more",
        [
            "Settings · Company · Billing",
            "Users & invites",
            "Permissions",
            "Integrations · Audit · Export",
            "Assistant · Alerts",
        ],
    ),
]
for i, (h, items) in enumerate(cols):
    x = Inches(0.35 + i * 3.25)
    add_rect(s, x, Inches(1.3), Inches(3.1), Inches(5.4), LIGHT)
    add_text(s, x + Inches(0.15), Inches(1.5), Inches(2.8), Inches(0.4), h, size=15, bold=True, color=NAVY)
    add_bullets(s, x + Inches(0.15), Inches(2.1), Inches(2.8), Inches(4.3), items, size=12)

# ── Inventory modules ────────────────────────────────────────────────
s = new_slide()
section_bar(s, "Inventory — Stock home", "Value overview, search, scan, checkout, adjust")
dual_viewport(
    s,
    "web_inventory_stock.png",
    "mobile_inventory_stock.png",
    "Web — Stock (Hyderabad Building Materials)",
    "Mobile — Stock + bottom tabs",
)

s = new_slide()
section_bar(s, "Inventory — Materials & Parties", "SKU catalog + customers/vendors")
img(s, "web_inventory_materials.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_inventory_parties.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Materials  ·  Right: Parties",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Inventory — Quotes & Sales", "Quote → sell / POS-style checkout path")
img(s, "web_inventory_quotes.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_inventory_sales.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Quotes  ·  Right: Sales",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Inventory — Warehouse & Procurement", "Transfers, staging, purchase side")
img(s, "web_inventory_warehouse.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_inventory_procurement.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Warehouse  ·  Right: Procurement",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Inventory — Invoices & Bills", "AR (sales invoices) and AP (purchase bills)")
img(s, "web_inventory_invoices.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_inventory_bills.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Invoices  ·  Right: Bills",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Inventory — Reports & Settings", "Ops insights + business profile / users")
img(s, "web_inventory_reports.png", Inches(0.35), Inches(1.2), Inches(6.3))
img(s, "web_inventory_settings.png", Inches(6.85), Inches(1.2), Inches(6.1))
add_text(
    s,
    Inches(0.35),
    Inches(6.55),
    Inches(12.5),
    Inches(0.35),
    "Left: Reports  ·  Right: Settings (profile, vertical, team)",
    size=12,
    color=MUTED,
)

s = new_slide()
section_bar(s, "Inventory on mobile", "Bottom tabs for field / counter use")
mobiles = [
    ("mobile_inventory_stock.png", "Stock"),
    ("mobile_inventory_sales.png", "Sales"),
    ("mobile_inventory_warehouse.png", "Warehouse"),
    ("mobile_inventory_invoices.png", "Invoices"),
]
for i, (name, label) in enumerate(mobiles):
    x = Inches(0.45 + i * 3.2)
    img(s, name, x, Inches(1.25), Inches(2.9))
    add_text(s, x, Inches(6.6), Inches(2.9), Inches(0.3), label, size=12, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

s = new_slide()
section_bar(s, "Inventory — function map", "Shell tabs + supporting screens")
add_bullets(
    s,
    Inches(0.6),
    Inches(1.4),
    Inches(6),
    Inches(5.5),
    [
        "Stock — on-hand, value, adjust, checkout, barcode scan",
        "Materials — SKU master, units, prices",
        "Parties — customers & vendors",
        "Quotes — customer quotations (events wording where vertical=EVENTS)",
        "Sales — sell / checkout flows",
        "Warehouse — multi-location, transfers",
        "Procurement — purchasing / inbound",
        "Invoices — sales AR with filters",
        "Bills — purchase AP with filters",
        "Reports — inventory analytics",
        "Settings — business profile, vertical, language, team",
        "Notifications + AI assistant FAB",
    ],
    size=15,
)
add_rect(s, Inches(7.2), Inches(1.4), Inches(5.6), Inches(5.2), LIGHT)
add_text(s, Inches(7.5), Inches(1.65), Inches(5.1), Inches(0.4), "Business profiles (demo)", size=15, bold=True, color=NAVY)
add_bullets(
    s,
    Inches(7.5),
    Inches(2.2),
    Inches(5.1),
    Inches(4),
    [
        "Material supplier",
        "Retail / Kirana (GENERAL + FEFO)",
        "Wholesale / cash & carry",
        "Distributor / stockist",
        "Trader",
        "Equipment dealer",
        "Events & lighting (EVENTS vertical)",
        "General business",
    ],
    size=14,
)

# ── Walkthrough + close ──────────────────────────────────────────────
s = new_slide()
section_bar(s, "Suggested team walkthrough", "30–40 minutes")
steps = [
    ("1", "Platform", "Log in as admin@buildflow.com — open Companies & Tickets."),
    ("2", "Construction", "owner@reddyconst.com — Dashboard → Projects → open NH-45 sample."),
    ("3", "Estimate path", "Proposals / Estimation → show BOQ & permissions."),
    ("4", "Site ops", "Planning + Reports + Assistant FAB."),
    ("5", "Finance", "Accounting + Reports Hub."),
    ("6", "Admin", "Settings → Users & Role permissions."),
    ("7", "Inventory", "owner@hydmaterials.com — Stock → Sales → Invoices."),
    ("8", "Verticals", "Optional: kirana-demo (FEFO) or luminalighting (events)."),
]
for i, (n, title, body) in enumerate(steps):
    col = i % 2
    row = i // 2
    x = Inches(0.4 + col * 6.45)
    y = Inches(1.25 + row * 1.4)
    add_rect(s, x, y, Inches(6.2), Inches(1.25), LIGHT)
    add_rect(s, x, y, Inches(0.55), Inches(1.25), NAVY)
    add_text(
        s,
        x + Inches(0.08),
        y + Inches(0.4),
        Inches(0.45),
        Inches(0.4),
        n,
        size=18,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    add_text(s, x + Inches(0.7), y + Inches(0.2), Inches(5.3), Inches(0.35), title, size=15, bold=True, color=NAVY)
    add_text(s, x + Inches(0.7), y + Inches(0.55), Inches(5.3), Inches(0.55), body, size=12, color=SLATE)

s = new_slide()
add_rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
add_rect(s, Inches(0), Inches(6.6), Inches(13.333), Inches(0.12), AMBER)
add_text(s, Inches(0.8), Inches(2.3), Inches(11.5), Inches(0.7), "Questions & next steps", size=40, bold=True, color=WHITE)
add_text(
    s,
    Inches(0.8),
    Inches(3.2),
    Inches(11),
    Inches(1.5),
    "Screenshots live under docs/team-ppt/screenshots/\n"
    "Re-capture anytime: node docs/team-ppt/capture-screens.mjs\n"
    "Credentials: docs/DEMO_CREDENTIALS.txt",
    size=16,
    color=RGBColor(0xE2, 0xE8, 0xF0),
)
add_text(
    s,
    Inches(0.8),
    Inches(5.3),
    Inches(11),
    Inches(0.4),
    "BuildFlow v2.0 — built for the field, designed for the boardroom.",
    size=14,
    color=AMBER,
)

total = len(slides_meta)
for i, slide in enumerate(slides_meta):
    if i == 0 or i == total - 1:
        continue
    footer(slide, i + 1, total)

prs.save(OUT_PPTX)
print("Wrote", OUT_PPTX)
print("Slides", total)
